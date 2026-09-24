import io
import json
from flask import Blueprint, render_template, request, send_file, flash, redirect, url_for
from flask_login import login_required
from app.models import HasilPenilaian, LembarJawaban, Siswa, MataPelajaran
from app.extensions import db
from app.routes.siswa import KELAS_LIST

laporan_bp = Blueprint('laporan', __name__, url_prefix='/laporan')


@laporan_bp.route('/')
@login_required
def index():
    mapel_id   = request.args.get('mapel_id', type=int)
    kelas      = request.args.get('kelas', '').strip()
    page       = request.args.get('page', 1, type=int)

    mapel_list  = MataPelajaran.query.order_by(MataPelajaran.nama).all()
    kelas_list  = KELAS_LIST

    query = (
        HasilPenilaian.query
        .join(HasilPenilaian.lembar_jawaban)
        .join(LembarJawaban.siswa)
        .join(LembarJawaban.mata_pelajaran)
    )

    if mapel_id:
        query = query.filter(LembarJawaban.mapel_id == mapel_id)
    if kelas:
        query = query.filter(Siswa.kelas == kelas)

    hasil_list = query.order_by(HasilPenilaian.created_at.desc()).paginate(page=page, per_page=15, error_out=False)

    return render_template(
        'laporan/index.html',
        title='Laporan Penilaian',
        hasil_list=hasil_list,
        mapel_list=mapel_list,
        kelas_list=kelas_list,
        selected_mapel=mapel_id,
        selected_kelas=kelas
    )


@laporan_bp.route('/detail/<int:id>')
@login_required
def detail(id):
    hasil = HasilPenilaian.query.get_or_404(id)
    detail_dict = {}
    if hasil.detail_jawaban:
        try:
            detail_dict = json.loads(hasil.detail_jawaban)
        except (json.JSONDecodeError, TypeError):
            detail_dict = {}

    return render_template('laporan/detail.html', title='Detail Penilaian', hasil=hasil, detail=detail_dict)


@laporan_bp.route('/export/excel')
@login_required
def export_excel():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        flash('Library openpyxl tidak ditemukan. Install dengan: pip install openpyxl', 'danger')
        return redirect(url_for('laporan.index'))

    mapel_id = request.args.get('mapel_id', type=int)
    kelas    = request.args.get('kelas', '').strip()

    query = (
        HasilPenilaian.query
        .join(HasilPenilaian.lembar_jawaban)
        .join(LembarJawaban.siswa)
        .join(LembarJawaban.mata_pelajaran)
    )
    if mapel_id:
        query = query.filter(LembarJawaban.mapel_id == mapel_id)
    if kelas:
        query = query.filter(Siswa.kelas == kelas)

    hasil_list = query.order_by(HasilPenilaian.created_at.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = 'Laporan Penilaian'

    # Header style
    header_fill = PatternFill(start_color='6C63FF', end_color='6C63FF', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')
    center      = Alignment(horizontal='center')

    headers = ['No', 'Nama Siswa', 'NIS', 'Kelas', 'Mata Pelajaran', 'Skor', 'Status', 'Tanggal']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill   = header_fill
        cell.font   = header_font
        cell.alignment = center

    for i, hasil in enumerate(hasil_list, 1):
        lembar = hasil.lembar_jawaban
        siswa  = lembar.siswa
        mapel  = lembar.mata_pelajaran
        ws.append([
            i,
            siswa.nama,
            siswa.nis,
            siswa.kelas,
            mapel.nama,
            hasil.skor or '-',
            lembar.status,
            hasil.created_at.strftime('%d/%m/%Y %H:%M')
        ])

    # Auto column width
    for col in ws.columns:
        max_len = max((len(str(cell.value or '')) for cell in col), default=0)
        ws.column_dimensions[col[0].column_letter].width = max_len + 4

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='laporan_penilaian.xlsx'
    )
